from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, A4
from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph,
    Spacer
)
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.enums import TA_CENTER
from datetime import datetime


def exportar_excel(rows, columns, filename):
    """
    Exporta una lista de datos a un archivo Excel.
    """

    wb = Workbook()
    ws = wb.active
    ws.title = "Datos"

    # Encabezados
    for col, nombre in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col, value=nombre)

        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    # Datos
    for row_index, row in enumerate(rows, start=2):

        for col_index, value in enumerate(row, start=1):

            if value is not None:
                value = str(value)

            ws.cell(
                row=row_index,
                column=col_index,
                value=value
            )

    # Ajustar ancho de columnas
    for column in ws.columns:

        max_length = 0
        column_letter = column[0].column_letter

        for cell in column:

            if cell.value is not None:
                max_length = max(
                    max_length,
                    len(str(cell.value))
                )

        ws.column_dimensions[column_letter].width = min(
            max_length + 2,
            40
        )

    wb.save(filename)


def exportar_pdf(rows, columns, filename, titulo):
    """
    Exporta una lista de datos a un PDF con formato profesional.
    """

    document = SimpleDocTemplate(
        filename,
        pagesize=landscape(A4),
        rightMargin=25,
        leftMargin=25,
        topMargin=25,
        bottomMargin=25
    )

    styles = getSampleStyleSheet()

    titulo_style = styles["Title"]
    titulo_style.alignment = TA_CENTER

    elementos = []

    elementos.append(
        Paragraph(titulo, titulo_style)
    )

    elementos.append(
        Spacer(1, 15)
    )

    # Convertir los datos a texto
    table_data = [columns]

    for row in rows:

        table_data.append([
            "" if value is None else str(value)
            for value in row
        ])

    tabla = Table(
        table_data,
        repeatRows=1
    )

    tabla.setStyle(
        TableStyle([
            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.HexColor("#2196F3")
            ),
            (
                "TEXTCOLOR",
                (0, 0),
                (-1, 0),
                colors.white
            ),
            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold"
            ),
            (
                "ALIGN",
                (0, 0),
                (-1, 0),
                "CENTER"
            ),
            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.5,
                colors.grey
            ),
            (
                "VALIGN",
                (0, 0),
                (-1, -1),
                "MIDDLE"
            ),
            (
                "FONTSIZE",
                (0, 0),
                (-1, -1),
                8
            ),
            (
                "ROWBACKGROUNDS",
                (0, 1),
                (-1, -1),
                [colors.white, colors.HexColor("#F2F2F2")]
            )
        ])
    )

    elementos.append(tabla)

    document.build(elementos)